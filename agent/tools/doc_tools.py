from __future__ import annotations

import json
from datetime import datetime
from email.message import EmailMessage
from pathlib import Path
from typing import Any, Dict, List

import pandas as pd

from agent.tools.base import ToolResult, fail, ok
from agent.templates.manager import TemplateManager
from agent.templates.generator import TemplateGenerator
from agent.templates.scanner import TemplateScanner

_template_manager = None
_template_generator = None

def get_template_manager():
    global _template_manager
    if _template_manager is None:
        _template_manager = TemplateManager()
    return _template_manager

def get_template_generator():
    global _template_generator
    if _template_generator is None:
        _template_generator = TemplateGenerator()
    return _template_generator



def _ensure_out_dir(output_dir: str | None) -> Path:
    path = Path(output_dir or "docs/parsed/reimburse_outputs").resolve()
    path.mkdir(parents=True, exist_ok=True)
    return path


# ======================
# 模板扫描，生成配置文件
# ======================
def scan_templates() -> ToolResult:
    """
    扫描所有模板文件，生成 templates_config.json 配置文件
    """
    try:
        scanner = TemplateScanner()
        config = scanner.generate_config_file()
        
        return ok(
            success=True,
            message=f"成功扫描 {len(config)} 个模板",
            template_count=len(config),
            config_file="data/templates/templates_config.json"
        )
    except Exception as e:
        return fail(f"扫描模板失败: {str(e)}")


def convert_xls_to_xlsx(xls_path: str, output_path: str = None) -> ToolResult:
    """
    将 .xls 文件转换为 .xlsx 文件
    
    Args:
        xls_path: .xls 文件路径
        output_path: 输出 .xlsx 文件路径，默认为同目录下同名 .xlsx 文件
    """
    try:
        scanner = TemplateScanner()
        converted_path = scanner.convert_xls_to_xlsx(xls_path, output_path)
        
        return ok(
            success=True,
            message="成功转换 .xls 文件为 .xlsx 文件",
            input_file=xls_path,
            output_file=converted_path
        )
    except Exception as e:
        return fail(f"转换 .xls 文件失败: {str(e)}")


def batch_convert_xls_to_xlsx() -> ToolResult:
    """
    批量转换目录中的 .xls 文件为 .xlsx 文件
    """
    try:
        scanner = TemplateScanner()
        converted_files = scanner.batch_convert_xls_to_xlsx()
        
        return ok(
            success=True,
            message=f"成功转换 {len(converted_files)} 个 .xls 文件",
            converted_count=len(converted_files),
            converted_files=converted_files
        )
    except Exception as e:
        return fail(f"批量转换 .xls 文件失败: {str(e)}")


# ======================
# 模板模糊匹配，确定需要的输入字段
# ======================
def search_templates_by_keyword(keyword: str) -> ToolResult:
    """
    基于关键词模糊匹配模板，并返回需要的输入字段
    从 collection_manifest.json 中搜索合适模板，然后从 templates_config 中获取占位符信息
    """
    try:
        manager = get_template_manager()
        
        # 首先从 collection_manifest.json 中搜索模板
        collection_templates = []
        collection_manifest_path = manager.templates_dir / "collection_manifest.json"
        if collection_manifest_path.exists():
            import json
            with open(collection_manifest_path, 'r', encoding='utf-8') as f:
                collection_info = json.load(f)
                templates = collection_info.get('templates', [])
                
                # 转换查询为小写，进行不区分大小写的匹配
                query_lower = keyword.lower()
                
                for template in templates:
                    template_name = template.get('target_name', '')
                    original_path = template.get('original_path', '')
                    parent_folders = template.get('parent_folders', [])
                    
                    # 检查模板名称是否匹配
                    if query_lower in template_name.lower():
                        collection_templates.append({
                            "name": template_name,
                            "original_path": original_path,
                            "parent_folders": parent_folders
                        })
                    # 检查原始路径是否匹配
                    elif query_lower in original_path.lower():
                        collection_templates.append({
                            "name": template_name,
                            "original_path": original_path,
                            "parent_folders": parent_folders
                        })
                    # 检查父文件夹是否匹配
                    elif any(query_lower in folder.lower() for folder in parent_folders):
                        collection_templates.append({
                            "name": template_name,
                            "original_path": original_path,
                            "parent_folders": parent_folders
                        })
        
        # 如果从 collection 中没有找到，使用原来的搜索方法
        if not collection_templates:
            collection_templates = manager.search_templates(keyword)
        
        # 提取每个模板的需要的输入字段
        template_info = []
        for template in collection_templates:
            template_name = template.get("name")
            
            if not template_name:
                continue
            
            # 从模板文件名中提取模板名称（移除扩展名）
            template_base_name = template_name
            if template_name:
                # 移除文件扩展名
                import os
                template_base_name = os.path.splitext(template_name)[0]
            
            # 从配置文件中获取模板信息
            # 先尝试使用原始模板名称
            template_config = manager.config.get(template_name, {})
            # 如果没有找到，尝试使用移除扩展名的模板名称
            if not template_config:
                template_config = manager.config.get(template_base_name, {})
            
            placeholders = template_config.get("placeholders", [])
            field_mapping = template_config.get("field_mapping", {})
            
            # 提取需要的参数类型
            required_fields = []
            for placeholder in placeholders:
                if isinstance(placeholder, str):
                    # 清理占位符格式，提取字段名
                    field_name = placeholder.strip('{}[]_')
                    required_fields.append({
                        "name": field_name,
                        "type": "string",  # 默认类型
                        "required": True
                    })
            
            info = {
                "name": template.get("name"),
                "type": template.get("type"),
                "placeholders": placeholders,
                "field_mapping": field_mapping,
                "required_fields": required_fields,
                "detailed_placeholders": [
                    {
                        "placeholder": placeholder,
                        "field_name": placeholder.strip('{}[]_'),
                        "mapped_field": field_mapping.get(placeholder.strip('{}[]_'), placeholder.strip('{}[]_'))
                    }
                    for placeholder in placeholders if isinstance(placeholder, str)
                ],
                "valid": template.get("valid", False),
                "original_path": template.get("original_path", ""),
                "parent_folders": template.get("parent_folders", [])
            }
            template_info.append(info)
        
        return ok(templates=template_info)
    except Exception as e:
        return fail(str(e))


# ======================
# 生成 Word（严格按你模板结构，不破坏、不新增表格）
# ======================
def generate_word_doc(activity: Dict[str, Any], invoices: List[Dict[str, Any]], output_dir: str | None = None, template_name: str = None) -> ToolResult:
    try:
        # 尝试使用新的模板管理系统
        generator = get_template_generator()
        manager = get_template_manager()
        
        # 如果没有指定模板，使用默认模板
        if not template_name:
            template_name = "学生活动经费使用情况.docx"
        
        # 从配置文件中获取模板信息
        # 先尝试使用原始模板名称
        template_config = manager.config.get(template_name, {})
        # 如果没有找到，尝试使用移除扩展名的模板名称
        if not template_config:
            import os
            template_base_name = os.path.splitext(template_name)[0]
            template_config = manager.config.get(template_base_name, {})
        # 如果还是没有找到，且模板是 .xls 格式，尝试使用 .xlsx 格式的名称
        if not template_config and template_name.lower().endswith('.xls'):
            xlsx_name = template_name[:-4] + '.xlsx'  # 将 .xls 替换为 .xlsx
            template_config = manager.config.get(xlsx_name, {})
        
        placeholders = template_config.get("placeholders", [])
        field_mapping = template_config.get("field_mapping", {})
        output_filename_pattern = template_config.get("output_filename_pattern", "activity_{timestamp}.docx")
        
        # 动态准备数据，根据模板的占位符
        data = {}
        for placeholder in placeholders:
            if isinstance(placeholder, str):
                # 清理占位符格式，提取字段名
                field_name = placeholder.strip('{}[]_')
                # 从 activity 中获取对应字段的值
                # 先尝试使用 field_mapping 映射的字段名
                mapped_field = field_mapping.get(field_name, field_name)
                data[field_name] = activity.get(mapped_field, activity.get(field_name, ""))
        
        # 使用配置文件中的命名模式生成文件名
        out_dir = _ensure_out_dir(output_dir)
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        output_filename = output_filename_pattern.replace('{timestamp}', timestamp)
        output_path = str(out_dir / output_filename)
        
        # 控制台输出
        print(f"\n=== 生成 Word 文档 ===")
        print(f"模板名称: {template_name}")
        print(f"输出路径: {output_path}")
        print(f"占位符数量: {len(placeholders)}")
        print(f"准备的数据: {data}")
        
        # 确定要传递给生成器的模板文件名（如果是 .xls，转换为 .xlsx）
        generator_template_name = template_name
        if template_name.lower().endswith('.xls'):
            generator_template_name = template_name[:-4] + '.xlsx'
        
        result = generator.generate_from_template(
            template_name=generator_template_name,
            data=data,
            output_path=output_path
        )
        
        if result.get("success"):
            print(f"[成功] 生成成功: {result['output_path']}")
            return ok(word_path=result["output_path"])
        else:
            print(f"[失败] 生成失败: {result.get('error', 'Unknown error')}")
            # 备用：使用旧的方法
            return _generate_word_doc_legacy(activity, invoices, output_dir)
            
    except Exception as e:
        print(f"[错误] 发生错误: {str(e)}")
        # 出错时使用旧方法
        return _generate_word_doc_legacy(activity, invoices, output_dir)

def _generate_word_doc_legacy(activity: Dict[str, Any], invoices: List[Dict[str, Any]], output_dir: str | None = None) -> ToolResult:
    out_dir = _ensure_out_dir(output_dir)
    target = out_dir / f"activity_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
    
    try:
        from docx import Document
    except Exception:
        return fail("python-docx 不可用")

    # 读取你本地的官方模板（保持结构100%一致）
    template_path = Path("data/templates") / "学生活动经费使用情况.docx"
    if not template_path.exists():
        template_path = Path(__file__).parent / "学生活动经费使用情况说明模板.docx"
    if not template_path.exists():
        return fail(f"模板不存在: {template_path}")

    doc = Document(str(template_path))

    # 要替换的内容（只替换文字，不碰表格结构）
    replace_map = {
        "经办人姓名": activity.get("student_name", ""),
        "经办人联系方式": activity.get("contact", ""),
        "活动时间": activity.get("activity_time", ""),
        "活动地点": activity.get("location", ""),
        "参与人员": activity.get("participants", ""),
        "活动主要内容": activity.get("description", ""),
        "报销内容及金额": activity.get("expense_detail", ""),
    }

    # 替换表格内文字
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text.strip()
                if "经办人姓名" in text:
                    cell.text = replace_map["经办人姓名"]
                elif "经办人联系方式" in text:
                    cell.text = replace_map["经办人联系方式"]
                elif "活动时间" in text:
                    cell.text = replace_map["活动时间"]
                elif "活动地点" in text:
                    cell.text = replace_map["活动地点"]
                elif "参与人员" in text:
                    cell.text = replace_map["参与人员"]
                elif "活动主要内容" in text:
                    cell.text = replace_map["活动主要内容"]
                elif "报销内容及金额" in text:
                    cell.text = replace_map["报销内容及金额"]

    doc.save(str(target))
    return ok(word_path=str(target))


# ======================
# 生成 Excel（完全按你模板字段，不破坏结构）
# ======================
def generate_excel_sheet(invoices: List[Dict[str, Any]], activity: Dict[str, Any], output_dir: str | None = None, template_name: str = None) -> ToolResult:
    try:
        # 尝试使用新的模板管理系统
        generator = get_template_generator()
        manager = get_template_manager()
        
        # 如果没有指定模板，使用默认模板
        if not template_name:
            template_name = "预算表.xlsx"
        
        # 从配置文件中获取模板信息
        # 先尝试使用原始模板名称
        template_config = manager.config.get(template_name, {})
        # 如果没有找到，尝试使用移除扩展名的模板名称
        if not template_config:
            import os
            template_base_name = os.path.splitext(template_name)[0]
            template_config = manager.config.get(template_base_name, {})
        # 如果还是没有找到，且模板是 .xls 格式，尝试使用 .xlsx 格式的名称
        if not template_config and template_name.lower().endswith('.xls'):
            xlsx_name = template_name[:-4] + '.xlsx'  # 将 .xls 替换为 .xlsx
            template_config = manager.config.get(xlsx_name, {})
        
        placeholders = template_config.get("placeholders", [])
        field_mapping = template_config.get("field_mapping", {})
        output_filename_pattern = template_config.get("output_filename_pattern", "reimburse_{timestamp}.xlsx")
        
        # 准备数据
        data = {
            "invoices": invoices,
            "activity": activity
        }
        
        # 动态添加模板需要的字段
        for placeholder in placeholders:
            if isinstance(placeholder, str):
                # 清理占位符格式，提取字段名
                field_name = placeholder.strip('{}[]_')
                # 从 activity 中获取对应字段的值
                # 先尝试使用 field_mapping 映射的字段名
                mapped_field = field_mapping.get(field_name, field_name)
                data[field_name] = activity.get(mapped_field, activity.get(field_name, ""))
        
        # 使用配置文件中的命名模式生成文件名
        out_dir = _ensure_out_dir(output_dir)
        output_filename_pattern = template_config.get("output_filename_pattern", "reimburse_{timestamp}.xlsx")
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        output_filename = output_filename_pattern.replace('{timestamp}', timestamp)
        output_path = str(out_dir / output_filename)
        
        # 控制台输出
        print(f"\n=== 生成 Excel 文档 ===")
        print(f"模板名称: {template_name}")
        print(f"输出路径: {output_path}")
        print(f"占位符数量: {len(placeholders)}")
        print(f"发票数量: {len(invoices)}")
        print(f"准备的数据: {data}")
        
        # 确定要传递给生成器的模板文件名（如果是 .xls，转换为 .xlsx）
        generator_template_name = template_name
        if template_name.lower().endswith('.xls'):
            generator_template_name = template_name[:-4] + '.xlsx'
        
        result = generator.generate_from_template(
            template_name=generator_template_name,
            data=data,
            output_path=output_path
        )
        
        if result.get("success"):
            print(f"[成功] 生成成功: {result['output_path']}")
            return ok(excel_path=result["output_path"])
        else:
            print(f"[失败] 生成失败: {result.get('error', 'Unknown error')}")
            # 备用：使用旧的方法
            return _generate_excel_sheet_legacy(invoices, activity, output_dir)
            
    except Exception as e:
        print(f"[错误] 发生错误: {str(e)}")
        # 出错时使用旧方法
        return _generate_excel_sheet_legacy(invoices, activity, output_dir)

def _generate_excel_sheet_legacy(invoices: List[Dict[str, Any]], activity: Dict[str, Any], output_dir: str | None = None) -> ToolResult:
    out_dir = _ensure_out_dir(output_dir)
    target = out_dir / f"reimburse_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

    rows = []
    for inv in invoices:
        rows.append({
            "发票序号": inv.get("invoice_no", ""),
            "发票金额": inv.get("amount", 0.0),
            "发票日期": inv.get("invoice_date", ""),
            "发票内容": inv.get("content", ""),
            "具体活动名称": activity.get("activity_name", ""),
            "活动举办日期": activity.get("activity_date", ""),
            "归属（学生组织）": activity.get("org", ""),
            "经办同学": activity.get("student_name", ""),
            "学号": activity.get("student_id", "")
        })

    df = pd.DataFrame(rows)
    df.to_excel(target, index=False, sheet_name="Sheet3")
    return ok(excel_path=str(target))


# ======================
# 邮件相关（保持你原来的不变）
# ======================
def generate_email_draft(activity: Dict[str, Any], summary: Dict[str, Any], attachments: List[str]) -> ToolResult:
    subject = f"报销材料提交 - {activity.get('activity_date', '')}"
    body = (
        "老师您好，\n\n"
        f"现提交活动报销材料。活动地点：{activity.get('location', '')}。"
        f"报销总金额：{summary.get('total_amount', 0)} 元。\n"
        "附件包含活动说明与报销明细表。\n\n"
        "此致\n敬礼"
    )
    return ok(draft={"subject": subject, "body": body, "attachments": attachments})


def send_or_export_email(draft: Dict[str, Any], output_dir: str | None = None) -> ToolResult:
    out_dir = _ensure_out_dir(output_dir)
    target = out_dir / f"mail_{datetime.now().strftime('%Y%m%d_%H%M%S')}.eml"
    msg = EmailMessage()
    msg["Subject"] = draft.get("subject", "报销材料")
    msg["To"] = ""
    msg["From"] = ""
    payload = {
        "body": draft.get("body", ""),
        "attachments": draft.get("attachments", []),
    }
    msg.set_content(json.dumps(payload, ensure_ascii=False, indent=2))
    target.write_text(msg.as_string(), encoding="utf-8")
    return ok(sent=False, eml_path=str(target), fallback_used=True)